# -*- coding: utf-8 -*-
"""完成版シミュレーションExcelの生成。

シート構成:
  1. シミュレーション   — パラメータ変更で全体が再計算されるライブモデル（数式ベース）
  2. シナリオ比較       — 保守/標準/強気の12ヶ月予測と料金プラン別手残りマトリクス
  3. モンテカルロ       — P10/P50/P90レンジと粗利が月額を上回る確率
  4. 実績データ         — 2025/03〜2026/06の全実績とCVR/平均単価
  5. 前提・根拠         — パラメータ推定根拠・手法・注意事項
"""
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS = json.load(open(os.path.join(BASE, "output", "simulation_results.json"), encoding="utf-8"))

# ---- styles -------------------------------------------------------------
NAVY = "1F3864"
BLUE = "2A78D6"
LIGHT = "DCE6F1"
GRAY = "F2F2F2"
YEN = '¥#,##0'
PCT2 = '0.00%'
PCT4 = '0.000%'

f_title = Font(bold=True, size=14, color=NAVY)
f_h = Font(bold=True, color="FFFFFF")
f_b = Font(bold=True)
f_note = Font(size=9, color="666666")
fill_h = PatternFill("solid", fgColor=NAVY)
fill_sub = PatternFill("solid", fgColor=LIGHT)
fill_act = PatternFill("solid", fgColor=GRAY)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.font = f_h
        cell.fill = fill_h
        cell.alignment = center
        cell.border = border


ACTUALS = [a for a in RESULTS["actuals"] if a["month"] >= "2026-01"]  # 直近6ヶ月を表示
FC_MONTHS = [m["month"] for m in RESULTS["scenarios"]["base"]["months"]]
P = RESULTS["params"]

wb = Workbook()

# ======================================================================
# 前提・根拠 シート（先に作る: シミュレーションシートの数式が参照するため）
# ======================================================================
ws_p = wb.active
ws_p.title = "前提・根拠"
ws_p.sheet_view.showGridLines = False
ws_p["B2"] = "前提条件とパラメータ推定根拠"
ws_p["B2"].font = f_title

rows = [
    ("実測パラメータ（2025/09〜2026/06 プール推定）", None, None, None),
    ("CVR（購入率）", P["cvr"], PCT4, "注文27件 ÷ 訪問42,218（95%信頼区間 0.044%〜0.093%）"),
    ("平均購入単価", P["aov"], YEN, "売上¥381,810 ÷ 27件（月次実績は¥7,963〜¥23,196と分散大）"),
    ("粗利率", P["gross_margin"], PCT2, "粗利¥187,109 ÷ 売上¥381,810"),
    ("直近6ヶ月のトラフィック月次成長率", P["traffic_mom_recent"], PCT2, "2026/01: 5,241 → 2026/06: 5,797（幾何平均）"),
    ("KW順位からの流入ポテンシャル", 1.48, '0.00"倍"', "追跡506KW中155KWが200位以内。CTRカーブ推定で現状1,799クリック/月→順位1段改善で2,657クリック/月"),
]
r = 4
for label, val, fmt, note in rows:
    ws_p.cell(r, 2, label)
    if val is None:
        ws_p.cell(r, 2).font = f_b
        ws_p.cell(r, 2).fill = fill_sub
        ws_p.cell(r, 3).fill = fill_sub
        ws_p.cell(r, 4).fill = fill_sub
    else:
        c = ws_p.cell(r, 3, val)
        c.number_format = fmt
        c.font = f_b
        ws_p.cell(r, 4, note).font = f_note
    r += 1

# ---- シナリオパラメータ表（シミュレーションシートが INDEX/MATCH で参照）----
r += 1
ws_p.cell(r, 2, "シナリオ別パラメータ（シミュレーションシートのドロップダウンと連動）").font = f_b
ws_p.cell(r, 2).fill = fill_sub
for cc in range(3, 7):
    ws_p.cell(r, cc).fill = fill_sub
r += 1
SC_TABLE_START = r + 1
hdr = ["シナリオ", "訪問数成長率/月", "CVR", "平均購入単価", "根拠"]
for i, h in enumerate(hdr):
    ws_p.cell(r, 2 + i, h)
style_header(ws_p, r, 2, 6)
sc_rows = [
    ("保守", 0.01, 0.00045, 12000, "成長が鈍化しCVRは信頼区間の下限で推移"),
    ("標準", 0.025, 0.00065, 14000, "直近実績の成長率+施策効果、CVR・単価は実測値を維持"),
    ("強気", 0.05, 0.0009, 15000, "KW順位改善(+48%流入余地)とCVR改善施策が奏功、CVRは信頼区間上限"),
    ("カスタム", 0.03, 0.0007, 14000, "自由に編集してください"),
]
for i, (nm, g, cvr, aov, note) in enumerate(sc_rows):
    rr = SC_TABLE_START + i
    ws_p.cell(rr, 2, nm).font = f_b
    ws_p.cell(rr, 3, g).number_format = PCT2
    ws_p.cell(rr, 4, cvr).number_format = PCT4
    ws_p.cell(rr, 5, aov).number_format = YEN
    ws_p.cell(rr, 6, note).font = f_note
    for cc in range(2, 7):
        ws_p.cell(rr, cc).border = border
r = SC_TABLE_START + len(sc_rows) + 1

# ---- 料金プラン表 ----
ws_p.cell(r, 2, "料金プラン（シミュレーションシートのドロップダウンと連動）").font = f_b
ws_p.cell(r, 2).fill = fill_sub
for cc in range(3, 7):
    ws_p.cell(r, cc).fill = fill_sub
r += 1
PLAN_TABLE_START = r + 1
hdr = ["プラン", "月額固定費", "売上レベシェア", "粗利レベシェア", "備考"]
for i, h in enumerate(hdr):
    ws_p.cell(r, 2 + i, h)
style_header(ws_p, r, 2, 6)
plan_rows = [
    ("現行（月額55,000円）", 55000, 0.0, 0.0, "現行契約"),
    ("提案（月額88,000円+売上5%）", 88000, 0.05, 0.0, "20260629シート記載の提案プラン"),
    ("希望（粗利レベシェア50%）", 0, 0.0, 0.50, "クライアント希望案。固定費ゼロ"),
    ("カスタム", 55000, 0.0, 0.0, "自由に編集してください"),
]
for i, (nm, fx, rs, gs, note) in enumerate(plan_rows):
    rr = PLAN_TABLE_START + i
    ws_p.cell(rr, 2, nm).font = f_b
    ws_p.cell(rr, 3, fx).number_format = YEN
    ws_p.cell(rr, 4, rs).number_format = PCT2
    ws_p.cell(rr, 5, gs).number_format = PCT2
    ws_p.cell(rr, 6, note).font = f_note
    for cc in range(2, 7):
        ws_p.cell(rr, cc).border = border
r = PLAN_TABLE_START + len(plan_rows) + 1

# ---- 手法・注意 ----
notes = [
    "【手法】",
    "・予測式: 訪問数[t] = 訪問数[t-1] × (1+成長率) / 注文数 = 訪問数 × CVR / 売上 = 注文数 × 平均単価 / 粗利 = 売上 × 粗利率",
    "・モンテカルロ: 注文数〜Poisson(訪問数×CVR)、注文単価〜対数正規（実測月次単価にフィット）、各月2万回試行",
    "・注文数が月3〜5件と少なくCVR推定の不確実性が大きいため、点推定ではなくP10〜P90のレンジで併記",
    "",
    "【元シミュレーションからの主な変更点】",
    "・CVR想定 0.10%→0.15% は実測95%信頼区間（0.044%〜0.093%）を大きく超えるため、実測ベースに補正",
    "・平均単価 ¥13,000 → 実測 ¥14,141 ベースに更新（¥14,000）",
    "・粗利率 50% → 実測 49.0%",
    "・トラフィック成長 110%/月 は直近実測（+2.0%/月）と乖離が大きいため、シナリオ別に 1%/2.5%/5% に補正",
    "",
    "【注意事項】",
    "※本数値は現時点の実績（注文27件）をもとに算出した概算シミュレーションです。",
    "※注文数のサンプルが少ないため、実際の売上は月次でP10〜P90レンジ程度の振れが想定されます。",
    "※実際の売上推移や成果は、運営状況・季節性・外部環境等の影響を受ける可能性があります。",
    "※今後の方向性や検討材料の一つとしてご参照いただくことを目的とした資料です。",
]
for line in notes:
    ws_p.cell(r, 2, line)
    if line.startswith("【"):
        ws_p.cell(r, 2).font = f_b
    else:
        ws_p.cell(r, 2).font = f_note
    r += 1

ws_p.column_dimensions["B"].width = 34
for col in "CDE":
    ws_p.column_dimensions[col].width = 16
ws_p.column_dimensions["F"].width = 70

# ======================================================================
# シミュレーション シート（数式ベースのライブモデル）
# ======================================================================
ws = wb.create_sheet("シミュレーション")
ws.sheet_view.showGridLines = False
ws["B2"] = "収益シミュレーション — armtechstore.jp"
ws["B2"].font = f_title
ws["B3"] = f"作成日: {RESULTS['generated']} ／ 下のシナリオ・料金プランを切り替えると全体が再計算されます"
ws["B3"].font = f_note

# パラメータブロック
ws["B5"] = "シナリオ選択"
ws["B5"].font = f_b
ws["C5"] = "標準"
ws["C5"].fill = fill_sub
ws["C5"].border = border
dv = DataValidation(type="list", formula1='"保守,標準,強気,カスタム"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(ws["C5"])

ws["B6"] = "料金プラン選択"
ws["B6"].font = f_b
ws["C6"] = "現行（月額55,000円）"
ws["C6"].fill = fill_sub
ws["C6"].border = border
plan_names = ",".join(p[0] for p in plan_rows)
dv2 = DataValidation(type="list", formula1=f'"{plan_names}"', allow_blank=False)
ws.add_data_validation(dv2)
dv2.add(ws["C6"])

sc_rng = f"'前提・根拠'!$B${SC_TABLE_START}:$B${SC_TABLE_START+3}"
ws["E5"] = "訪問数成長率/月"
ws["F5"] = f"=INDEX('前提・根拠'!$C${SC_TABLE_START}:$C${SC_TABLE_START+3},MATCH($C$5,{sc_rng},0))"
ws["F5"].number_format = PCT2
ws["E6"] = "CVR"
ws["F6"] = f"=INDEX('前提・根拠'!$D${SC_TABLE_START}:$D${SC_TABLE_START+3},MATCH($C$5,{sc_rng},0))"
ws["F6"].number_format = PCT4
ws["E7"] = "平均購入単価"
ws["F7"] = f"=INDEX('前提・根拠'!$E${SC_TABLE_START}:$E${SC_TABLE_START+3},MATCH($C$5,{sc_rng},0))"
ws["F7"].number_format = YEN
ws["E8"] = "粗利率"
ws["F8"] = P["gross_margin"]
ws["F8"].number_format = PCT2

pl_rng = f"'前提・根拠'!$B${PLAN_TABLE_START}:$B${PLAN_TABLE_START+3}"
ws["H5"] = "月額固定費"
ws["I5"] = f"=INDEX('前提・根拠'!$C${PLAN_TABLE_START}:$C${PLAN_TABLE_START+3},MATCH($C$6,{pl_rng},0))"
ws["I5"].number_format = YEN
ws["H6"] = "売上レベシェア"
ws["I6"] = f"=INDEX('前提・根拠'!$D${PLAN_TABLE_START}:$D${PLAN_TABLE_START+3},MATCH($C$6,{pl_rng},0))"
ws["I6"].number_format = PCT2
ws["H7"] = "粗利レベシェア"
ws["I7"] = f"=INDEX('前提・根拠'!$E${PLAN_TABLE_START}:$E${PLAN_TABLE_START+3},MATCH($C$6,{pl_rng},0))"
ws["I7"].number_format = PCT2
for addr in ("E5", "E6", "E7", "E8", "H5", "H6", "H7"):
    ws[addr].font = f_note

# グリッド
GRID = 10
n_act = len(ACTUALS)
n_all = n_act + len(FC_MONTHS)
ws.cell(GRID, 2, "")
ws.cell(GRID - 1, 3, "実績").font = f_b
ws.cell(GRID - 1, 3 + n_act, "予測").font = f_b
for i in range(n_all):
    c = 3 + i
    m = (ACTUALS + [{"month": x} for x in [mm["month"] for mm in RESULTS["scenarios"]["base"]["months"]]])[i]["month"]
    cell = ws.cell(GRID, c, datetime.strptime(m + "-01", "%Y-%m-%d"))
    cell.number_format = 'yyyy/mm'
style_header(ws, GRID, 2, 2 + n_all)
ws.cell(GRID, 2, "項目")

labels = ["サイト訪問数", "注文数", "売上", "粗利", "CVR", "平均購入単価", "SEO費用", "手残り（粗利−費用）", "累計手残り"]
for i, lb in enumerate(labels):
    cell = ws.cell(GRID + 1 + i, 2, lb)
    cell.font = f_b
    cell.border = border

R_VIS, R_ORD, R_REV, R_GP, R_CVR, R_AOV, R_FEE, R_NET, R_CUM = range(GRID + 1, GRID + 10)

# 実績列
for i, a in enumerate(ACTUALS):
    c = 3 + i
    cl = get_column_letter(c)
    ws.cell(R_VIS, c, a["visits"]).number_format = '#,##0'
    ws.cell(R_ORD, c, a["orders"]).number_format = '#,##0'
    ws.cell(R_REV, c, a["revenue"]).number_format = YEN
    ws.cell(R_GP, c, a["gp"]).number_format = YEN
    ws.cell(R_CVR, c, f"=IFERROR({cl}{R_ORD}/{cl}{R_VIS},0)").number_format = PCT4
    ws.cell(R_AOV, c, f"=IFERROR({cl}{R_REV}/{cl}{R_ORD},0)").number_format = YEN
    ws.cell(R_FEE, c, 55000).number_format = YEN
    ws.cell(R_NET, c, f"={cl}{R_GP}-{cl}{R_FEE}").number_format = YEN
    if i == 0:
        ws.cell(R_CUM, c, f"={cl}{R_NET}").number_format = YEN
    else:
        pv = get_column_letter(c - 1)
        ws.cell(R_CUM, c, f"={pv}{R_CUM}+{cl}{R_NET}").number_format = YEN
    for rr in range(R_VIS, R_CUM + 1):
        ws.cell(rr, c).fill = fill_act
        ws.cell(rr, c).border = border

# 予測列（数式）
for i in range(len(FC_MONTHS)):
    c = 3 + n_act + i
    cl = get_column_letter(c)
    pv = get_column_letter(c - 1)
    ws.cell(R_VIS, c, f"=ROUND({pv}{R_VIS}*(1+$F$5),0)").number_format = '#,##0'
    ws.cell(R_ORD, c, f"={cl}{R_VIS}*$F$6").number_format = '0.0'
    ws.cell(R_REV, c, f"={cl}{R_ORD}*$F$7").number_format = YEN
    ws.cell(R_GP, c, f"={cl}{R_REV}*$F$8").number_format = YEN
    ws.cell(R_CVR, c, "=$F$6").number_format = PCT4
    ws.cell(R_AOV, c, "=$F$7").number_format = YEN
    ws.cell(R_FEE, c, f"=$I$5+{cl}{R_REV}*$I$6+{cl}{R_GP}*$I$7").number_format = YEN
    ws.cell(R_NET, c, f"={cl}{R_GP}-{cl}{R_FEE}").number_format = YEN
    ws.cell(R_CUM, c, f"={pv}{R_CUM}+{cl}{R_NET}").number_format = YEN
    for rr in range(R_VIS, R_CUM + 1):
        ws.cell(rr, c).border = border

# 12ヶ月合計列
c_tot = 3 + n_all
cl_first = get_column_letter(3 + n_act)
cl_last = get_column_letter(2 + n_all)
ws.cell(GRID, c_tot, "予測12ヶ月計")
style_header(ws, GRID, c_tot, c_tot)
for rr, fmt in [(R_ORD, '0.0'), (R_REV, YEN), (R_GP, YEN), (R_FEE, YEN), (R_NET, YEN)]:
    ws.cell(rr, c_tot, f"=SUM({cl_first}{rr}:{cl_last}{rr})").number_format = fmt
    ws.cell(rr, c_tot).font = f_b
    ws.cell(rr, c_tot).border = border

ws.column_dimensions["B"].width = 22
for i in range(3, c_tot + 1):
    ws.column_dimensions[get_column_letter(i)].width = 11

# 注釈
rr = R_CUM + 2
ws.cell(rr, 2, "※実績列（灰色）は確定値。予測列はシナリオパラメータに連動する数式です。"
             "レンジ（P10〜P90）は「モンテカルロ」シートをご参照ください。").font = f_note

# グラフ: 粗利と費用の推移
chart = LineChart()
chart.title = "粗利とSEO費用の推移（実績+予測）"
chart.style = 2
chart.height = 8
chart.width = 24
chart.y_axis.numFmt = YEN
chart.y_axis.title = "円"
data_gp = Reference(ws, min_col=2, min_row=R_GP, max_col=2 + n_all, max_row=R_GP)
data_fee = Reference(ws, min_col=2, min_row=R_FEE, max_col=2 + n_all, max_row=R_FEE)
cats = Reference(ws, min_col=3, min_row=GRID, max_col=2 + n_all, max_row=GRID)
chart.add_data(data_gp, titles_from_data=True, from_rows=True)
chart.add_data(data_fee, titles_from_data=True, from_rows=True)
chart.set_categories(cats)
s0, s1 = chart.series
s0.graphicalProperties.line.solidFill = "2A78D6"
s0.graphicalProperties.line.width = 20000
s1.graphicalProperties.line.solidFill = "E34948"
s1.graphicalProperties.line.width = 20000
s1.graphicalProperties.line.dashStyle = "dash"
ws.add_chart(chart, f"B{rr + 2}")

# ======================================================================
# シナリオ比較 シート
# ======================================================================
ws_s = wb.create_sheet("シナリオ比較")
ws_s.sheet_view.showGridLines = False
ws_s["B2"] = "シナリオ比較（2026/07〜2027/06）"
ws_s["B2"].font = f_title

r = 4
for key in ("conservative", "base", "optimistic"):
    sc = RESULTS["scenarios"][key]
    a = sc["assumptions"]
    ws_s.cell(r, 2, f"【{sc['label']}】 成長率{a['traffic_growth_mom']*100:.1f}%/月・CVR {a['cvr']*100:.3f}%・単価¥{a['aov']:,}  —  {a['note']}").font = f_b
    ws_s.cell(r, 2).fill = fill_sub
    for cc in range(3, 16):
        ws_s.cell(r, cc).fill = fill_sub
    r += 1
    ws_s.cell(r, 2, "項目")
    for i, m in enumerate(sc["months"]):
        ws_s.cell(r, 3 + i, datetime.strptime(m["month"] + "-01", "%Y-%m-%d")).number_format = 'yyyy/mm'
    ws_s.cell(r, 15, "12ヶ月計")
    style_header(ws_s, r, 2, 15)
    r += 1
    for label, kk, fmt in [("サイト訪問数", "visits", '#,##0'), ("注文数", "orders", '0.0'),
                           ("売上", "revenue", YEN), ("粗利", "gp", YEN)]:
        ws_s.cell(r, 2, label).font = f_b
        ws_s.cell(r, 2).border = border
        for i, m in enumerate(sc["months"]):
            cell = ws_s.cell(r, 3 + i, m[kk])
            cell.number_format = fmt
            cell.border = border
        if kk in ("revenue", "gp"):
            tot = ws_s.cell(r, 15, sc["total_revenue"] if kk == "revenue" else sc["total_gp"])
            tot.number_format = YEN
            tot.font = f_b
            tot.border = border
        r += 1
    r += 1

# 料金プラン別 手残りマトリクス
ws_s.cell(r, 2, "料金プラン別 クライアント手残り（12ヶ月累計 = 粗利 − SEO費用）").font = f_b
ws_s.cell(r, 2).fill = fill_sub
for cc in range(3, 6):
    ws_s.cell(r, cc).fill = fill_sub
r += 1
ws_s.cell(r, 2, "プラン ＼ シナリオ")
for i, key in enumerate(("conservative", "base", "optimistic")):
    ws_s.cell(r, 3 + i, RESULTS["scenarios"][key]["label"])
style_header(ws_s, r, 2, 5)
r += 1
for pk in ("current", "proposed", "gp_share_50"):
    ws_s.cell(r, 2, RESULTS["scenarios"]["base"]["fee_plans"][pk]["label"]).font = f_b
    ws_s.cell(r, 2).border = border
    for i, key in enumerate(("conservative", "base", "optimistic")):
        v = RESULTS["scenarios"][key]["fee_plans"][pk]["total_client_net"]
        cell = ws_s.cell(r, 3 + i, v)
        cell.number_format = YEN
        cell.border = border
        if v < 0:
            cell.font = Font(color="C00000")
    r += 1
ws_s.cell(r + 1, 2, "※粗利レベシェア50%案は費用が粗利に連動するため、クライアント手残りは常に粗利の50%（マイナスにならない）。"
                    "固定費型プランは粗利が費用を下回る月は持ち出しになります。").font = f_note

ws_s.column_dimensions["B"].width = 30
for i in range(3, 16):
    ws_s.column_dimensions[get_column_letter(i)].width = 11

# ======================================================================
# モンテカルロ シート
# ======================================================================
ws_m = wb.create_sheet("モンテカルロ")
ws_m.sheet_view.showGridLines = False
ws_m["B2"] = "モンテカルロ・シミュレーション（各月2万回試行）"
ws_m["B2"].font = f_title
ws_m["B3"] = "注文数〜Poisson(訪問数×CVR)、注文単価〜対数正規（実測フィット）。売上・粗利の振れ幅を示します。"
ws_m["B3"].font = f_note

r = 5
for key in ("conservative", "base", "optimistic"):
    sc = RESULTS["scenarios"][key]
    ws_m.cell(r, 2, f"【{sc['label']}】").font = f_b
    ws_m.cell(r, 2).fill = fill_sub
    for cc in range(3, 15):
        ws_m.cell(r, cc).fill = fill_sub
    r += 1
    ws_m.cell(r, 2, "項目")
    for i, m in enumerate(sc["monte_carlo"]):
        ws_m.cell(r, 3 + i, datetime.strptime(m["month"] + "-01", "%Y-%m-%d")).number_format = 'yyyy/mm'
    style_header(ws_m, r, 2, 14)
    r += 1
    for label, kk, fmt in [
        ("売上 P90（上振れ）", "revenue_p90", YEN),
        ("売上 P50（中央値）", "revenue_p50", YEN),
        ("売上 P10（下振れ）", "revenue_p10", YEN),
        ("粗利 P50", "gp_p50", YEN),
        ("粗利≧55,000円の確率", "prob_gp_over_55k", '0%'),
    ]:
        ws_m.cell(r, 2, label).font = f_b
        ws_m.cell(r, 2).border = border
        for i, m in enumerate(sc["monte_carlo"]):
            cell = ws_m.cell(r, 3 + i, m[kk])
            cell.number_format = fmt
            cell.border = border
        r += 1
    r += 1
ws_m.column_dimensions["B"].width = 24
for i in range(3, 15):
    ws_m.column_dimensions[get_column_letter(i)].width = 11

# ======================================================================
# 実績データ シート
# ======================================================================
ws_a = wb.create_sheet("実績データ")
ws_a.sheet_view.showGridLines = False
ws_a["B2"] = "月次実績（2025/03〜2026/06）"
ws_a["B2"].font = f_title
r = 4
hdr = ["年月", "サイト訪問数", "注文数", "売上", "粗利", "CVR", "平均購入単価", "粗利率"]
for i, h in enumerate(hdr):
    ws_a.cell(r, 2 + i, h)
style_header(ws_a, r, 2, 9)
r += 1
for a in RESULTS["actuals"]:
    ws_a.cell(r, 2, datetime.strptime(a["month"] + "-01", "%Y-%m-%d")).number_format = 'yyyy/mm'
    ws_a.cell(r, 3, a["visits"]).number_format = '#,##0'
    ws_a.cell(r, 4, a["orders"]).number_format = '#,##0'
    ws_a.cell(r, 5, a["revenue"]).number_format = YEN
    ws_a.cell(r, 6, a["gp"]).number_format = YEN
    ws_a.cell(r, 7, f"=IFERROR(D{r}/C{r},0)").number_format = PCT4
    ws_a.cell(r, 8, f"=IFERROR(E{r}/D{r},0)").number_format = YEN
    ws_a.cell(r, 9, f"=IFERROR(F{r}/E{r},0)").number_format = PCT2
    for cc in range(2, 10):
        ws_a.cell(r, cc).border = border
    r += 1
# 合計行
ws_a.cell(r, 2, "合計/平均").font = f_b
for cc, fml, fmt in [(3, f"=SUM(C5:C{r-1})", '#,##0'), (4, f"=SUM(D5:D{r-1})", '#,##0'),
                     (5, f"=SUM(E5:E{r-1})", YEN), (6, f"=SUM(F5:F{r-1})", YEN),
                     (7, f"=D{r}/C{r}", PCT4), (8, f"=E{r}/D{r}", YEN), (9, f"=F{r}/E{r}", PCT2)]:
    cell = ws_a.cell(r, cc, fml)
    cell.number_format = fmt
    cell.font = f_b
    cell.border = border
ws_a.cell(r, 2).border = border
for col, w in zip("BCDEFGHI", [12, 14, 10, 12, 12, 10, 14, 10]):
    ws_a.column_dimensions[col].width = w

# シート順序: シミュレーションを先頭へ
wb.move_sheet("シミュレーション", -(len(wb.sheetnames) - 1))
wb.move_sheet("前提・根拠", 4)

out = os.path.join(BASE, "output", "収益シミュレーション_armtechstore_2026-07.xlsx")
wb.save(out)
print("written:", out)
